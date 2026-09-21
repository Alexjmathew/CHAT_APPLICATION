from flask import (
    Flask,
    render_template,
    request,
    jsonify,
    session,
    redirect,
    url_for
)

from werkzeug.utils import secure_filename

from werkzeug.security import (
    generate_password_hash,
    check_password_hash
)

from pypdf import PdfReader
from docx import Document

from openai import OpenAI

from functools import wraps

import sqlite3
import os
import uuid
import json
import csv


# ============================================================
# APP
# ============================================================

app = Flask(__name__)


app.secret_key = os.environ.get(
    "FLASK_SECRET_KEY",
    "change-this-secret-key"
)


# ============================================================
# LM STUDIO
# ============================================================

client = OpenAI(
    base_url="http://127.0.0.1:1234/v1",
    api_key="lm-studio"
)


# Normal text model
TEXT_MODEL = "qwen/qwen3-coder-30b"


# Vision model
#
# IMPORTANT:
# This must exactly match the model ID shown
# in LM Studio.
#
VISION_MODEL = "gemma-4-12b"


# ============================================================
# DATABASE
# ============================================================

DATABASE = "chat.db"


# ============================================================
# FILE UPLOADS
# ============================================================

UPLOAD_FOLDER = "uploads"


ALLOWED_EXTENSIONS = {
    "pdf",
    "txt",
    "md",
    "csv",
    "json",
    "docx"
}


# Maximum uploaded file size
MAX_FILE_SIZE = 20 * 1024 * 1024  # 20 MB


# Maximum extracted document text
MAX_DOCUMENT_CHARS = 100000


os.makedirs(
    UPLOAD_FOLDER,
    exist_ok=True
)


app.config["UPLOAD_FOLDER"] = UPLOAD_FOLDER

app.config["MAX_CONTENT_LENGTH"] = MAX_FILE_SIZE


# ============================================================
# ALLOWED FILE
# ============================================================

def allowed_file(filename):

    if not filename:
        return False

    if "." not in filename:
        return False

    extension = filename.rsplit(
        ".",
        1
    )[1].lower()

    return extension in ALLOWED_EXTENSIONS


# ============================================================
# DATABASE CONNECTION
# ============================================================

def get_db():

    connection = sqlite3.connect(
        DATABASE
    )

    connection.row_factory = sqlite3.Row

    return connection


# ============================================================
# DATABASE INITIALIZATION
# ============================================================

def init_db():

    db = get_db()


    db.execute(
        "PRAGMA foreign_keys = ON"
    )


    db.executescript(
        """

        CREATE TABLE IF NOT EXISTS users (

            id INTEGER PRIMARY KEY AUTOINCREMENT,

            username TEXT UNIQUE NOT NULL,

            password TEXT NOT NULL,

            created_at
                TIMESTAMP DEFAULT CURRENT_TIMESTAMP

        );


        CREATE TABLE IF NOT EXISTS conversations (

            id INTEGER PRIMARY KEY AUTOINCREMENT,

            user_id INTEGER NOT NULL,

            title TEXT NOT NULL
                DEFAULT 'New Chat',

            created_at
                TIMESTAMP DEFAULT CURRENT_TIMESTAMP,

            updated_at
                TIMESTAMP DEFAULT CURRENT_TIMESTAMP,

            FOREIGN KEY(user_id)
                REFERENCES users(id)
                ON DELETE CASCADE

        );


        CREATE TABLE IF NOT EXISTS messages (

            id INTEGER PRIMARY KEY AUTOINCREMENT,

            conversation_id INTEGER NOT NULL,

            role TEXT NOT NULL,

            content TEXT NOT NULL,

            image_path TEXT,

            created_at
                TIMESTAMP DEFAULT CURRENT_TIMESTAMP,

            FOREIGN KEY(conversation_id)
                REFERENCES conversations(id)
                ON DELETE CASCADE

        );

        """
    )


    db.commit()

    db.close()


init_db()


# ============================================================
# AUTH DECORATOR
# ============================================================

def login_required(function):

    @wraps(function)
    def wrapper(*args, **kwargs):

        if not session.get("user_id"):

            return redirect(
                url_for("login")
            )

        return function(
            *args,
            **kwargs
        )

    return wrapper


# ============================================================
# FILE TEXT EXTRACTION
# ============================================================

def extract_pdf_text(file_path):

    reader = PdfReader(
        file_path
    )


    pages = []


    for page in reader.pages:

        try:

            text = page.extract_text()

        except Exception as error:

            print(
                "PDF page extraction error:",
                repr(error)
            )

            text = None


        if text:

            pages.append(
                text
            )


    return "\n\n".join(
        pages
    )


# ============================================================
# DOCX
# ============================================================

def extract_docx_text(file_path):

    document = Document(
        file_path
    )


    paragraphs = []


    for paragraph in document.paragraphs:

        text = paragraph.text.strip()


        if text:

            paragraphs.append(
                text
            )


    # Also extract table contents
    for table in document.tables:

        for row in table.rows:

            cells = []


            for cell in row.cells:

                cell_text = (
                    cell.text
                    .strip()
                )


                if cell_text:

                    cells.append(
                        cell_text
                    )


            if cells:

                paragraphs.append(
                    " | ".join(cells)
                )


    return "\n".join(
        paragraphs
    )


# ============================================================
# TEXT / MARKDOWN
# ============================================================

def extract_text_file(file_path):

    with open(
        file_path,
        "r",
        encoding="utf-8",
        errors="ignore"
    ) as file:

        return file.read()


# ============================================================
# CSV
# ============================================================

def extract_csv_text(file_path):

    rows = []


    with open(
        file_path,
        "r",
        encoding="utf-8",
        errors="ignore",
        newline=""
    ) as file:

        reader = csv.reader(
            file
        )


        for row in reader:

            rows.append(
                " | ".join(row)
            )


    return "\n".join(
        rows
    )


# ============================================================
# JSON
# ============================================================

def extract_json_text(file_path):

    with open(
        file_path,
        "r",
        encoding="utf-8",
        errors="ignore"
    ) as file:

        data = json.load(
            file
        )


    return json.dumps(
        data,
        indent=2,
        ensure_ascii=False
    )


# ============================================================
# GENERAL FILE EXTRACTION
# ============================================================

def extract_file_text(
    file_path,
    extension
):

    extension = extension.lower()


    if extension == "pdf":

        return extract_pdf_text(
            file_path
        )


    if extension == "docx":

        return extract_docx_text(
            file_path
        )


    if extension in {
        "txt",
        "md"
    }:

        return extract_text_file(
            file_path
        )


    if extension == "csv":

        return extract_csv_text(
            file_path
        )


    if extension == "json":

        return extract_json_text(
            file_path
        )


    raise ValueError(
        f"Unsupported file type: {extension}"
    )


# ============================================================
# REGISTER
# ============================================================

@app.route(
    "/register",
    methods=["GET", "POST"]
)
def register():

    if session.get("user_id"):

        return redirect(
            url_for("home")
        )


    if request.method == "POST":

        username = request.form.get(
            "username",
            ""
        ).strip()


        password = request.form.get(
            "password",
            ""
        )


        if len(username) < 3:

            return render_template(
                "register.html",
                error=(
                    "Username must contain "
                    "at least 3 characters."
                )
            )


        if len(password) < 6:

            return render_template(
                "register.html",
                error=(
                    "Password must contain "
                    "at least 6 characters."
                )
            )


        password_hash = (
            generate_password_hash(
                password
            )
        )


        db = get_db()


        try:

            cursor = db.execute(
                """
                INSERT INTO users
                (
                    username,
                    password
                )
                VALUES (?, ?)
                """,
                (
                    username,
                    password_hash
                )
            )


            db.commit()


            user_id = cursor.lastrowid


        except sqlite3.IntegrityError:

            db.close()


            return render_template(
                "register.html",
                error="Username already exists."
            )


        db.close()


        session["user_id"] = user_id

        session["username"] = username


        return redirect(
            url_for("home")
        )


    return render_template(
        "register.html"
    )


# ============================================================
# LOGIN
# ============================================================

@app.route(
    "/login",
    methods=["GET", "POST"]
)
def login():

    if session.get("user_id"):

        return redirect(
            url_for("home")
        )


    if request.method == "POST":

        username = request.form.get(
            "username",
            ""
        ).strip()


        password = request.form.get(
            "password",
            ""
        )


        db = get_db()


        user = db.execute(
            """
            SELECT *
            FROM users
            WHERE username = ?
            """,
            (username,)
        ).fetchone()


        db.close()


        if (
            user
            and check_password_hash(
                user["password"],
                password
            )
        ):

            session["user_id"] = (
                user["id"]
            )

            session["username"] = (
                user["username"]
            )


            return redirect(
                url_for("home")
            )


        return render_template(
            "login.html",
            error=(
                "Invalid username or password."
            )
        )


    return render_template(
        "login.html"
    )


# ============================================================
# LOGOUT
# ============================================================

@app.route("/logout")
def logout():

    session.clear()

    return redirect(
        url_for("login")
    )


# ============================================================
# HOME
# ============================================================

@app.route("/")
@login_required
def home():

    return render_template(
        "index.html",
        username=session.get(
            "username"
        )
    )


# ============================================================
# GET CONVERSATIONS
# ============================================================

@app.route(
    "/api/conversations"
)
@login_required
def get_conversations():

    db = get_db()


    conversations = db.execute(
        """
        SELECT
            id,
            title,
            created_at,
            updated_at
        FROM conversations
        WHERE user_id = ?
        ORDER BY updated_at DESC
        """,
        (
            session["user_id"],
        )
    ).fetchall()


    db.close()


    return jsonify([
        dict(conversation)
        for conversation in conversations
    ])


# ============================================================
# CREATE CONVERSATION
# ============================================================

@app.route(
    "/api/conversations",
    methods=["POST"]
)
@login_required
def create_conversation():

    data = (
        request.get_json(
            silent=True
        )
        or {}
    )


    title = data.get(
        "title",
        "New Chat"
    )


    if not isinstance(
        title,
        str
    ):

        title = "New Chat"


    title = title.strip()


    if not title:

        title = "New Chat"


    db = get_db()


    cursor = db.execute(
        """
        INSERT INTO conversations
        (
            user_id,
            title
        )
        VALUES (?, ?)
        """,
        (
            session["user_id"],
            title
        )
    )


    db.commit()


    conversation_id = (
        cursor.lastrowid
    )


    db.close()


    return jsonify({
        "id": conversation_id,
        "title": title
    })


# ============================================================
# DELETE CONVERSATION
# ============================================================

@app.route(
    "/api/conversations/<int:conversation_id>",
    methods=["DELETE"]
)
@login_required
def delete_conversation(
    conversation_id
):

    db = get_db()


    db.execute(
        """
        DELETE FROM conversations
        WHERE id = ?
        AND user_id = ?
        """,
        (
            conversation_id,
            session["user_id"]
        )
    )


    db.commit()

    db.close()


    return jsonify({
        "success": True
    })


# ============================================================
# UPLOAD DOCUMENT
# ============================================================

@app.route(
    "/api/upload",
    methods=["POST"]
)
@login_required
def upload_file():

    if "file" not in request.files:

        return jsonify({
            "error": "No file was uploaded."
        }), 400


    file = request.files["file"]


    if not file.filename:

        return jsonify({
            "error": "No file selected."
        }), 400


    if not allowed_file(
        file.filename
    ):

        return jsonify({
            "error": (
                "Unsupported file type. "
                "Allowed: "
                "PDF, TXT, MD, CSV, JSON, DOCX."
            )
        }), 400


    original_filename = (
        secure_filename(
            file.filename
        )
    )


    if not original_filename:

        return jsonify({
            "error": "Invalid filename."
        }), 400


    extension = (
        original_filename
        .rsplit(".", 1)[1]
        .lower()
    )


    stored_filename = (
        str(uuid.uuid4())
        + "."
        + extension
    )


    file_path = os.path.join(
        app.config["UPLOAD_FOLDER"],
        stored_filename
    )


    try:

        # ----------------------------------------------------
        # Save uploaded file
        # ----------------------------------------------------

        file.save(
            file_path
        )


        # ----------------------------------------------------
        # Extract text
        # ----------------------------------------------------

        extracted_text = (
            extract_file_text(
                file_path,
                extension
            )
        )


        extracted_text = (
            extracted_text.strip()
        )


        if not extracted_text:

            try:

                os.remove(
                    file_path
                )

            except OSError:

                pass


            return jsonify({
                "error": (
                    "No readable text was "
                    "found in this file."
                )
            }), 400


        # ----------------------------------------------------
        # Limit document size
        # ----------------------------------------------------

        truncated = False


        if len(extracted_text) > (
            MAX_DOCUMENT_CHARS
        ):

            extracted_text = (
                extracted_text[
                    :MAX_DOCUMENT_CHARS
                ]
            )

            truncated = True


        # ----------------------------------------------------
        # Return document
        # ----------------------------------------------------

        return jsonify({

            "success": True,

            "filename":
                original_filename,

            "file_type":
                extension,

            "content":
                extracted_text,

            "truncated":
                truncated

        })


    except Exception as error:

        print(
            "File extraction error:",
            repr(error)
        )


        try:

            if os.path.exists(
                file_path
            ):

                os.remove(
                    file_path
                )

        except OSError:

            pass


        return jsonify({
            "error": (
                "Could not process "
                "the uploaded file: "
                + str(error)
            )
        }), 500


# ============================================================
# GET MESSAGES
# ============================================================

@app.route(
    "/api/conversations/<int:conversation_id>"
)
@login_required
def get_messages(
    conversation_id
):

    db = get_db()


    conversation = db.execute(
        """
        SELECT *
        FROM conversations
        WHERE id = ?
        AND user_id = ?
        """,
        (
            conversation_id,
            session["user_id"]
        )
    ).fetchone()


    if not conversation:

        db.close()


        return jsonify({
            "error": (
                "Conversation not found."
            )
        }), 404


    messages = db.execute(
        """
        SELECT
            id,
            role,
            content,
            image_path,
            created_at
        FROM messages
        WHERE conversation_id = ?
        ORDER BY id ASC
        """,
        (
            conversation_id,
        )
    ).fetchall()


    db.close()


    return jsonify({

        "conversation":
            dict(conversation),

        "messages": [
            dict(message)
            for message in messages
        ]

    })


# ============================================================
# CHAT
# ============================================================

@app.route(
    "/chat",
    methods=["POST"]
)
@login_required
def chat():

    data = (
        request.get_json(
            silent=True
        )
        or {}
    )


    # --------------------------------------------------------
    # Request data
    # --------------------------------------------------------

    conversation_id = data.get(
        "conversation_id"
    )


    user_message = data.get(
        "message",
        ""
    )


    if not isinstance(
        user_message,
        str
    ):

        user_message = ""


    user_message = (
        user_message.strip()
    )


    image_data = data.get(
        "image"
    )


    document_content = data.get(
        "document_content",
        ""
    )


    if not isinstance(
        document_content,
        str
    ):

        document_content = ""


    document_content = (
        document_content.strip()
    )


    document_name = data.get(
        "document_name",
        ""
    )


    if not isinstance(
        document_name,
        str
    ):

        document_name = ""


    document_name = (
        document_name.strip()
    )


    # --------------------------------------------------------
    # Validate conversation
    # --------------------------------------------------------

    if not conversation_id:

        return jsonify({
            "error": (
                "Conversation ID is required."
            )
        }), 400


    # --------------------------------------------------------
    # Validate message
    # --------------------------------------------------------

    if (
        not user_message
        and not image_data
        and not document_content
    ):

        return jsonify({
            "error": (
                "Message, image, or "
                "document is required."
            )
        }), 400


    db = get_db()


    try:

        # ----------------------------------------------------
        # Verify conversation belongs to user
        # ----------------------------------------------------

        conversation = db.execute(
            """
            SELECT *
            FROM conversations
            WHERE id = ?
            AND user_id = ?
            """,
            (
                conversation_id,
                session["user_id"]
            )
        ).fetchone()


        if not conversation:

            db.close()


            return jsonify({
                "error": (
                    "Conversation not found."
                )
            }), 404


        # ----------------------------------------------------
        # Get previous messages
        # ----------------------------------------------------

        previous_messages = (
            db.execute(
                """
                SELECT
                    role,
                    content,
                    image_path
                FROM messages
                WHERE conversation_id = ?
                ORDER BY id ASC
                """,
                (
                    conversation_id,
                )
            ).fetchall()
        )


        # ----------------------------------------------------
        # Determine what user message to save
        # ----------------------------------------------------

        if user_message:

            saved_user_content = (
                user_message
            )

        elif document_content:

            saved_user_content = (
                f"[Document: "
                f"{document_name or 'Uploaded document'}]"
            )

        elif image_data:

            saved_user_content = (
                "[Image]"
            )

        else:

            saved_user_content = (
                "[Attachment]"
            )


        # ----------------------------------------------------
        # Save user message
        # ----------------------------------------------------

        db.execute(
            """
            INSERT INTO messages
            (
                conversation_id,
                role,
                content
            )
            VALUES (?, ?, ?)
            """,
            (
                conversation_id,
                "user",
                saved_user_content
            )
        )


        # ----------------------------------------------------
        # Update conversation timestamp
        # ----------------------------------------------------

        db.execute(
            """
            UPDATE conversations
            SET updated_at =
                CURRENT_TIMESTAMP
            WHERE id = ?
            """,
            (
                conversation_id,
            )
        )


        db.commit()


        # ----------------------------------------------------
        # Build LM Studio messages
        # ----------------------------------------------------

        api_messages = [

            {
                "role": "system",

                "content": (
                    "You are a helpful, friendly "
                    "general-purpose AI assistant. "
                    "Answer clearly and naturally. "
                    "Use Markdown when useful. "
                    "When writing code, use fenced "
                    "Markdown code blocks with the "
                    "correct language. "
                    "Do not invent information. "
                    "If information is provided in "
                    "a document, use that document "
                    "as the primary source."
                )
            }

        ]


        # ----------------------------------------------------
        # Previous conversation
        # ----------------------------------------------------

        for message in previous_messages:

            role = message["role"]

            content = message["content"]


            if role not in {
                "user",
                "assistant",
                "system"
            }:

                continue


            api_messages.append({

                "role": role,

                "content": content

            })


        # ====================================================
        # IMAGE REQUEST
        # ====================================================

        if image_data:

            selected_model = (
                VISION_MODEL
            )


            if (
                not selected_model
                or selected_model
                == "your-vision-model"
            ):

                db.close()


                return jsonify({
                    "error": (
                        "Vision model is not "
                        "configured. Set "
                        "VISION_MODEL in app.py "
                        "to the exact vision model "
                        "name shown by LM Studio."
                    )
                }), 400


            vision_content = []


            if user_message:

                vision_content.append({

                    "type": "text",

                    "text": user_message

                })

            else:

                vision_content.append({

                    "type": "text",

                    "text": (
                        "Describe and analyze "
                        "this image."
                    )

                })


            vision_content.append({

                "type": "image_url",

                "image_url": {

                    "url": image_data

                }

            })


            api_messages.append({

                "role": "user",

                "content": vision_content

            })


        # ====================================================
        # DOCUMENT REQUEST
        # ====================================================

        elif document_content:

            selected_model = (
                TEXT_MODEL
            )


            document_prompt = f"""
The user has uploaded a document.

Filename:
{document_name or "Uploaded document"}

Document contents:
---------------- BEGIN DOCUMENT ----------------

{document_content}

----------------- END DOCUMENT -----------------

Instructions:

1. Use the document as the primary source.
2. Answer the user's question clearly.
3. Do not invent facts that are not supported
   by the document.
4. If the answer cannot be found in the document,
   say that clearly.
5. You may use your general knowledge when useful,
   but distinguish it from information found in
   the document.

User's question:

{user_message or "Analyze and summarize this document."}
"""


            api_messages.append({

                "role": "user",

                "content": document_prompt

            })


        # ====================================================
        # NORMAL TEXT REQUEST
        # ====================================================

        else:

            selected_model = (
                TEXT_MODEL
            )


            api_messages.append({

                "role": "user",

                "content": user_message

            })


        # ====================================================
        # ASK LM STUDIO
        # ====================================================

        print(
            "Sending request to LM Studio:",
            selected_model
        )

        response = client.chat.completions.create(
            model=selected_model,
            messages=api_messages,
            temperature=0.7,
            max_tokens=2048
        )


        # ----------------------------------------------------
        # Get response
        # ----------------------------------------------------

        if (
            not response.choices
        ):

            raise RuntimeError(
                "LM Studio returned no choices."
            )


        answer = (
            response
            .choices[0]
            .message
            .content
        )


        if not answer:

            answer = (
                "The model returned "
                "an empty response."
            )


        # ----------------------------------------------------
        # Save assistant message
        # ----------------------------------------------------

        db.execute(
            """
            INSERT INTO messages
            (
                conversation_id,
                role,
                content
            )
            VALUES (?, ?, ?)
            """,
            (
                conversation_id,
                "assistant",
                answer
            )
        )


        # ----------------------------------------------------
        # Generate conversation title
        # ----------------------------------------------------

        if (
            conversation["title"]
            == "New Chat"
        ):

            if user_message:

                title = (
                    user_message[:45]
                )

            elif document_name:

                title = (
                    document_name[:45]
                )

            elif image_data:

                title = "Image conversation"

            else:

                title = "New Chat"


            db.execute(
                """
                UPDATE conversations
                SET
                    title = ?,
                    updated_at =
                        CURRENT_TIMESTAMP
                WHERE id = ?
                """,
                (
                    title,
                    conversation_id
                )
            )


        # ----------------------------------------------------
        # Commit
        # ----------------------------------------------------

        db.commit()

        db.close()


        # ----------------------------------------------------
        # Return
        # ----------------------------------------------------

        return jsonify({

            "message": answer,

            "model": selected_model

        })


    except Exception as error:

        print(
            "LM Studio error:",
            repr(error)
        )


        try:

            db.rollback()

        except Exception:

            pass


        try:

            db.close()

        except Exception:

            pass


        return jsonify({

            "error":
                str(error)

        }), 500


# ============================================================
# FILE TOO LARGE
# ============================================================

@app.errorhandler(413)
def file_too_large(error):

    return jsonify({
        "error": (
            "File is too large. "
            "Maximum size is 20 MB."
        )
    }), 413


# ============================================================
# RUN
# ============================================================

if __name__ == "__main__":

    app.run(
        host="0.0.0.0",
        port=5000,
        debug=True
    )
